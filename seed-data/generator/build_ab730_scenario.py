# /// script
# dependencies = ["pyyaml", "jsonschema", "openpyxl", "python-docx"]
# ///

import argparse
import hashlib
import importlib.util
import io
import json
import re
import tempfile
import uuid
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips


HERE = Path(__file__).resolve().parent
BUILD_SCENARIO_PATH = HERE / "build_scenario.py"

EXPECTED_ROLE_KEYS = [
    "Admin",
    "Sponsor",
    "Planner",
    "Supply",
    "Quality",
    "Operations",
    "Engineering",
    "Analytics",
]


def _load_build_scenario_module():
    spec = importlib.util.spec_from_file_location("build_scenario", BUILD_SCENARIO_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _find_git_roots(path: Path) -> list[Path]:
    resolved = path.resolve()
    return [anc for anc in (resolved, *resolved.parents) if (anc / ".git").exists()]


def _assert_safe_private_staging_root(staging_root: Path, output_dir: Path) -> Path:
    resolved_staging = staging_root.resolve()
    resolved_output = output_dir.resolve()
    staging_chain = (resolved_staging, *resolved_staging.parents)
    output_chain = (resolved_output, *resolved_output.parents)

    if resolved_output in staging_chain or resolved_staging in output_chain:
        raise ValueError(
            f"Unsafe staging directory {resolved_staging}: overlaps with output directory {resolved_output}"
        )
    if resolved_staging.parent == resolved_output.parent:
        raise ValueError(
            f"Unsafe staging directory {resolved_staging}: cannot share parent with output directory {resolved_output}"
        )
    if _find_git_roots(resolved_staging):
        raise ValueError(
            f"Unsafe staging directory {resolved_staging}: private staging must never reside inside a Git working tree"
        )
    for git_root in [*_find_git_roots(resolved_output), *_find_git_roots(HERE)]:
        if git_root in staging_chain:
            raise ValueError(
                f"Unsafe staging directory {resolved_staging}: resides inside repository root {git_root}"
            )
    return resolved_staging


def _resolve_staging_root(output_dir: Path) -> Path:
    candidates: list[Path] = []
    if len(HERE.parents) > 2:
        session_files = HERE.parents[2] / "files"
        if session_files.exists() and session_files.is_dir():
            candidates.append(session_files / "ab730-generator-staging")

    temp_base = Path(tempfile.gettempdir()).resolve()
    if not temp_base.as_posix().startswith(("/tmp", "/var/tmp")):
        candidates.append(temp_base / "ab730-generator-staging")

    for candidate in candidates:
        try:
            return _assert_safe_private_staging_root(candidate, output_dir)
        except ValueError:
            continue

    raise ValueError(
        "Cannot resolve a safe private staging directory outside Git working trees and output directory"
    )


def _extract_roles_and_forbidden_tokens(identity_config: dict, role_keys: list[str]) -> tuple[dict, str, list[str]]:
    raw_roles = {}
    if isinstance(identity_config.get("roles"), dict):
        for role_key, info in identity_config["roles"].items():
            if isinstance(info, dict):
                raw_roles[role_key] = {
                    "upn": str(info.get("upn", "")).strip(),
                    "displayName": str(info.get("displayName", role_key)).strip() or role_key,
                }
    elif isinstance(identity_config.get("roleMap"), list):
        for entry in identity_config["roleMap"]:
            if isinstance(entry, dict) and "roleKey" in entry:
                rk = str(entry["roleKey"])
                raw_roles[rk] = {
                    "upn": str(entry.get("upn", "")).strip(),
                    "displayName": str(entry.get("displayName", rk)).strip() or rk,
                }
    else:
        raise ValueError("Identity config must provide 'roles' object or 'roleMap' list")

    missing = [rk for rk in role_keys if rk not in raw_roles]
    if missing:
        raise ValueError(f"Identity config is missing required roles: {', '.join(missing)}")

    domains = set()
    seen_upns = set()
    staged_roles = {}
    forbidden_tokens = [
        "moneyyu.com",
        "onmicrosoft.com",
    ]

    for secret_key in ("clientSecret", "tenantId", "clientId", "accessToken"):
        val = identity_config.get(secret_key)
        if isinstance(val, str) and val.strip() and not val.strip().startswith("<"):
            forbidden_tokens.append(val.strip())
    target_obj = identity_config.get("target")
    if isinstance(target_obj, dict):
        for secret_key in ("tenantId", "clientId"):
            val = target_obj.get(secret_key)
            if isinstance(val, str) and val.strip() and not val.strip().startswith("<"):
                forbidden_tokens.append(val.strip())

    admin_upn_entry = raw_roles["Admin"]["upn"]
    for top_admin_key in ("adminUpn", "demoUserUpn"):
        top_val = identity_config.get(top_admin_key)
        if isinstance(top_val, str) and top_val.strip():
            if top_val.strip().casefold() != admin_upn_entry.casefold():
                raise ValueError(f"Identity config {top_admin_key} does not match roles.Admin.upn")

    for rk in role_keys:
        entry = raw_roles[rk]
        upn = entry["upn"]
        if "@" not in upn or re.search(r"\s", upn):
            raise ValueError(f"Invalid UPN format for role {rk}")
        local_part, domain = upn.rsplit("@", 1)
        if not local_part or not domain or "." not in domain:
            raise ValueError(f"Invalid UPN domain for role {rk}")
        upn_cf = upn.casefold()
        if upn_cf in seen_upns:
            raise ValueError(f"Duplicate UPN assignment detected for role {rk}")
        seen_upns.add(upn_cf)
        domains.add(domain.casefold())
        staged_roles[rk] = {
            "upn": upn,
            "displayName": entry["displayName"],
        }
        if domain.casefold() != "example.invalid":
            forbidden_tokens.append(domain)
            forbidden_tokens.append(upn)
        if entry["displayName"] not in role_keys:
            forbidden_tokens.append(entry["displayName"])

    if len(domains) != 1:
        raise ValueError("All configured roles must belong to a single tenant domain")

    tenant_domain = next(iter(domains))
    return staged_roles, tenant_domain, sorted(set(t for t in forbidden_tokens if len(t) >= 3))


def validate_spec_and_inputs(spec: dict, identity_config: dict, date: str, output_dir: Path) -> tuple[dict, str, list[str]]:
    if output_dir.exists():
        existing_entries = [p.name for p in output_dir.iterdir()]
        if existing_entries:
            raise FileExistsError(
                f"Refusing to write into non-empty output directory {output_dir}: {existing_entries}"
            )

    if not re.fullmatch(r"\d{8}", str(date)):
        raise ValueError(f"Invalid --date {date!r}: expected YYYYMMDD")
    datetime.strptime(date, "%Y%m%d")

    if spec.get("courseCode") != "AB-730":
        raise ValueError("Spec courseCode must be AB-730")
    if spec.get("deliveryDate") != date:
        raise ValueError(f"Date {date} does not match spec deliveryDate {spec.get('deliveryDate')}")
    if spec.get("locale") != "zh-CN" or spec.get("currency") != "CNY":
        raise ValueError("Spec must declare locale zh-CN and currency CNY")
    if spec.get("adminIsBusinessRole") is not False:
        raise ValueError("Spec must enforce adminIsBusinessRole == False")

    role_keys = spec.get("roleKeys", [])
    if role_keys != EXPECTED_ROLE_KEYS:
        raise ValueError("Spec roleKeys do not match canonical 8 roles")

    office_files = spec.get("officeFiles", [])
    if len(office_files) != 8:
        raise ValueError("Spec must declare exactly 8 officeFiles")

    sp_docs = spec.get("sharePointDocuments", [])
    if sp_docs != [
        "06-Quality-Brief.docx",
        "07-Supplier-Response.docx",
        "08-Improvement-Meeting.docx",
    ]:
        raise ValueError("Spec sharePointDocuments must match the 3 canonical Word documents")

    email_rules = spec.get("emailRules", {})
    if "Admin" not in email_rules.get("ccEveryMessage", []):
        raise ValueError("emailRules.ccEveryMessage must include Admin")
    if email_rules.get("externalRecipientsAllowed") is not False:
        raise ValueError("emailRules.externalRecipientsAllowed must be False")

    if len(spec.get("emailThreads", [])) != 5:
        raise ValueError("Spec must declare 5 emailThreads")
    for thread in spec["emailThreads"]:
        if len(thread.get("authors", [])) != 3:
            raise ValueError("Every email thread must have 3 authors (1 send + 2 replies)")
        if len(thread.get("replyToRoles", [])) != 2:
            raise ValueError("Every email thread must have 2 replyToRoles entries")

    story_times = spec.get("teamStoryTimes", {})
    if int(story_times.get("dayOffset", 1)) > 0:
        raise ValueError("teamStoryTimes.dayOffset must not be in the future")
    if len(spec.get("teamChannels", [])) != 5:
        raise ValueError("Spec must declare 5 teamChannels")
    for ch in spec["teamChannels"]:
        if len(ch.get("topAuthors", [])) != 2 or not ch.get("firstPostReplyAuthor"):
            raise ValueError("Each team channel must have 2 topAuthors and 1 firstPostReplyAuthor")

    return _extract_roles_and_forbidden_tokens(identity_config, role_keys)


def _build_workbooks_spec(spec: dict) -> list[dict]:
    bom = spec["bom"]
    deliveries = spec["deliveries"]
    inspections = spec["inspections"]
    production = spec["production"]
    actions = spec["actions"]

    bom_rows = [
        [part_id, qty, base_cost, curr_cost, None, None, None]
        for part_id, qty, base_cost, curr_cost in bom["rows"]
    ]

    delivery_rows = []
    for po_id, part_id, sup_id, prom_date, act_date, planned_qty, recv_qty in deliveries["rows"]:
        if act_date is None:
            on_time_status = "Unknown"
        elif act_date <= prom_date:
            on_time_status = "OnTime"
        else:
            on_time_status = "Delayed"
        delivery_rows.append(
            [
                po_id,
                part_id,
                sup_id,
                prom_date,
                act_date,
                planned_qty,
                recv_qty,
                None,
                on_time_status,
                None,
            ]
        )

    quality_notes = {
        "Q-001": "常规进料抽检批次（不良率 2.00%）",
        "Q-002": "P-103 批次抽检异常（不良率 10.00%，触发 A-01）",
        "Q-003": "P-108 待检批次，抽检数 0 显式返回 N/A（D-02，非 0%）",
    }
    quality_rows = [
        [
            batch_id,
            part_id,
            inspected_qty,
            defective_qty,
            None,
            quality_notes.get(batch_id, "独立进料抽检范围"),
        ]
        for batch_id, part_id, inspected_qty, defective_qty in inspections["rows"]
    ]

    production_rows = [
        [
            line_id,
            planned_qty,
            actual_qty,
            None,
            None,
            "虚构局部总成排程缺口，与供应商交期缺口 150 件无直接因果换算，不自动推论停线",
        ]
        for line_id, planned_qty, actual_qty in production["rows"]
    ]

    action_titles = {
        "A-01": "排查 Q-002（P-103）批次不良率 10.00% 根因并复测",
        "A-02": "确认 PO-002（S-02）延迟到货与缺口 50 件补齐安排",
        "A-03": "提案：确认 PO-005（P-111）实际到货日与尾数 50 件追补方案（待确认）",
        "A-04": "核对 P-101 单价由 500 元上调至 550 元（单台 +100 元）技术变更依据",
        "A-05": "核实 LINE-01/LINE-02 合计 70 台排程缺口边界，明确非 150 件交期缺口直接导致",
        "A-06": "按加权口径 30/1100 = 2.73% 输出质量周报（禁止算术平均）",
    }
    action_rows = [
        [
            row["id"],
            row["source"],
            action_titles.get(row["id"], row["status"]),
            row["ownerRole"],
            "是" if row["ownerConfirmed"] else "否",
            row["dueDate"],
            row["status"],
        ]
        for row in actions
    ]

    return [
        {
            "filename": "01-BOM-Cost.xlsx",
            "disclaimerNote": (
                "本表为 AB-730 课程虚构训练资料（单位：人民币元 CNY），仅统计虚构局部总成"
                "（12 项零件 P-101..P-112），不代表整车 BOM 成本。基准合计 2,582 元，现行合计 2,682 元，差异 +100 元。"
            ),
            "sheets": [
                {
                    "name": "子系统BOM成本",
                    "table": "BOMCost",
                    "headers": [
                        "零件编号",
                        "单台用量",
                        "基准单价(元)",
                        "现行单价(元)",
                        "基准小计(元)",
                        "现行小计(元)",
                        "单台成本差异(元)",
                    ],
                    "rows": bom_rows,
                    "formulas": {
                        "基准小计(元)": "=B{row}*C{row}",
                        "现行小计(元)": "=B{row}*D{row}",
                        "单台成本差异(元)": "=F{row}-E{row}",
                    },
                    "numberFormats": {
                        "单台用量": "#,##0",
                        "基准单价(元)": "#,##0",
                        "现行单价(元)": "#,##0",
                        "基准小计(元)": "#,##0",
                        "现行小计(元)": "#,##0",
                        "单台成本差异(元)": "#,##0",
                    },
                    "summaryLabel": "局部总成合计",
                    "summaryFormulas": {
                        "基准小计(元)": "=SUM(E2:E13)",
                        "现行小计(元)": "=SUM(F2:F13)",
                        "单台成本差异(元)": "=SUM(G2:G13)",
                    },
                }
            ],
        },
        {
            "filename": "02-Supplier-Delivery.xlsx",
            "disclaimerNote": (
                "本表为 AB-730 课程虚构训练资料。五笔采购单计划交货合计 1,200 件，实收合计 1,050 件，交期缺口 150 件"
                "（PO-005 实际到货日为空，准时状态 Unknown，标记 D-01，不得默认视为准时）。"
            ),
            "sheets": [
                {
                    "name": "供应商交期追踪",
                    "table": "SupplierDelivery",
                    "headers": [
                        "采购单号",
                        "零件编号",
                        "供应商编号",
                        "承诺交期",
                        "实际到货日",
                        "计划数量",
                        "实收数量",
                        "缺口数量",
                        "准时状态",
                        "交期达成率",
                    ],
                    "rows": delivery_rows,
                    "dateColumns": ["承诺交期", "实际到货日"],
                    "formulas": {
                        "缺口数量": "=F{row}-G{row}",
                        "交期达成率": '=IF(F{row}=0,"N/A",G{row}/F{row})',
                    },
                    "numberFormats": {
                        "计划数量": "#,##0",
                        "实收数量": "#,##0",
                        "缺口数量": "#,##0",
                        "交期达成率": "0.00%",
                    },
                    "summaryLabel": "合计",
                    "summaryFormulas": {
                        "计划数量": "=SUM(F2:F6)",
                        "实收数量": "=SUM(G2:G6)",
                        "缺口数量": "=SUM(H2:H6)",
                        "交期达成率": '=IF(SUM(F2:F6)=0,"N/A",SUM(G2:G6)/SUM(F2:F6))',
                    },
                }
            ],
        },
        {
            "filename": "03-Quality-Inspection.xlsx",
            "disclaimerNote": (
                "本表为 AB-730 课程虚构训练资料。Q-001 至 Q-003 为独立生产/库存抽检范围，非 02-Supplier-Delivery.xlsx"
                " 五笔采购单普查映射。Q-003 抽检数为 0 返回 N/A（D-02）；加权不良率 30/1100 = 2.73%（D-04）。"
            ),
            "sheets": [
                {
                    "name": "进料抽检记录",
                    "table": "QualityInspection",
                    "headers": [
                        "抽检批号",
                        "零件编号",
                        "抽检数量",
                        "不良数量",
                        "不良率",
                        "范围与异常说明",
                    ],
                    "rows": quality_rows,
                    "formulas": {
                        "不良率": '=IF(C{row}=0,"N/A",D{row}/C{row})',
                    },
                    "numberFormats": {
                        "抽检数量": "#,##0",
                        "不良数量": "#,##0",
                        "不良率": "0.00%",
                    },
                    "summaryLabel": "加权合计",
                    "summaryFormulas": {
                        "抽检数量": "=SUM(C2:C4)",
                        "不良数量": "=SUM(D2:D4)",
                        "不良率": '=IF(SUM(C2:C4)=0,"N/A",SUM(D2:D4)/SUM(C2:C4))',
                    },
                }
            ],
        },
        {
            "filename": "04-Production-Impact.xlsx",
            "disclaimerNote": (
                "本表为 AB-730 课程虚构训练资料。计划产出合计 900 台，实际产出 830 台，虚构局部总成排程缺口 70 台；"
                "该 70 台缺口与供应商交期缺口 150 件无直接因果换算，不得自动推论停线决策。"
            ),
            "sheets": [
                {
                    "name": "产线排程影响",
                    "table": "ProductionImpact",
                    "headers": [
                        "产线编号",
                        "计划产出(台)",
                        "实际产出(台)",
                        "产出缺口(台)",
                        "达成率",
                        "排程说明与边界",
                    ],
                    "rows": production_rows,
                    "formulas": {
                        "产出缺口(台)": "=B{row}-C{row}",
                        "达成率": '=IF(B{row}=0,"N/A",C{row}/B{row})',
                    },
                    "numberFormats": {
                        "计划产出(台)": "#,##0",
                        "实际产出(台)": "#,##0",
                        "产出缺口(台)": "#,##0",
                        "达成率": "0.00%",
                    },
                    "summaryLabel": "合计",
                    "summaryFormulas": {
                        "计划产出(台)": "=SUM(B2:B3)",
                        "实际产出(台)": "=SUM(C2:C3)",
                        "产出缺口(台)": "=SUM(D2:D3)",
                        "达成率": '=IF(SUM(B2:B3)=0,"N/A",SUM(C2:C3)/SUM(B2:B3))',
                    },
                }
            ],
        },
        {
            "filename": "05-Action-Tracker.xlsx",
            "disclaimerNote": (
                "本表为 AB-730 课程虚构训练资料。共 6 项改善行动（A-01 至 A-06），其中 A-03（来源 PO-005，"
                "负责角色 Supply）为提案未确认（责任人已确认=否，目标完成日为空，状态 ProposalOnly，标记 D-03）。"
            ),
            "sheets": [
                {
                    "name": "改善行动追踪",
                    "table": "ImprovementActions",
                    "headers": [
                        "行动编号",
                        "来源单号",
                        "改善事项",
                        "负责角色",
                        "责任人已确认",
                        "目标完成日",
                        "状态",
                    ],
                    "rows": action_rows,
                    "dateColumns": ["目标完成日"],
                }
            ],
        },
    ]


def _build_emails_spec(spec: dict) -> dict:
    cc_roles = list(spec["emailRules"]["ccEveryMessage"])
    thread_bodies = {
        "【AB-730 20261015 Automotive】供应商交期与缺料风险": [
            "依据 02-Supplier-Delivery.xlsx 与 07-Supplier-Response.docx，五笔采购单计划交货 1,200 件、实收 1,050 件，总交期缺口 150 件（PO-002、PO-004、PO-005 各缺 50 件）。其中 PO-005（P-111，S-02）实收 150 件但实际到货日未核实（actualDate 为空，准时状态 Unknown，D-01）。",
            "收到。Operations 核对排程时不会将 PO-005 默认视为准时到货，也请 Supply 跟进 S-02 交期确认单。",
            "Supply 已记录：A-02（PO-002 补货确认，目标日 2026-10-16）已确认；但针对 PO-005 的 A-03 仍为 ProposalOnly（未确认责任人与截止日，D-03）。",
        ],
        "【AB-730 20261015 Automotive】BOM成本变动核对": [
            "依据 01-BOM-Cost.xlsx（表格 BOMCost），虚构局部总成 12 项零件（P-101..P-112）基准单台成本合计 2,582 元，现行单台成本合计 2,682 元，净差异 +100 元（单位：CNY，非整车 BOM）。",
            "Analytics 复核完毕：仅 P-101（单台用量 2 件）单价由 500 元调至 550 元（+100 元/台），其余 P-102..P-112 单价未变。",
            "Engineering 确认已登记行动项 A-04（2026-10-18 前完成 P-101 成本变动技术核验）。",
        ],
        "【AB-730 20261015 Automotive】零件批次质量异常": [
            "依据 03-Quality-Inspection.xlsx 与 06-Quality-Brief.docx，三批独立抽检数据为 Q-001（1000 检 20 不良，2.00%）、Q-002（100 检 10 不良，10.00%）、Q-003（0 检 0 不良，返回 N/A，D-02）。有效抽检合计 1,100 件、不良 30 件，加权不良率 2.73%（D-04）。",
            "Engineering 已会同 Quality 启动 Q-002（P-103）异常排查，注意汇总汇报时不得把 2.00% 与 10.00% 简单平均成 6.00%，且 Q-003 不得算作 0%。",
            "Quality 确认已建立行动项 A-01（来源 Q-002，目标日 2026-10-17），并在 06-Quality-Brief.docx 中注明抽检批次并非五笔采购单普查。",
        ],
        "【AB-730 20261015 Automotive】生产影响评估": [
            "依据 04-Production-Impact.xlsx 与 08-Improvement-Meeting.docx，LINE-01（500/450，缺口 50 台）与 LINE-02（400/380，缺口 20 台）合计虚构局部总成产出缺口 70 台。",
            "Planner 提醒：04-Production-Impact.xlsx 的 70 台局部总成缺口与 02-Supplier-Delivery.xlsx 的 150 件零件交期缺口单位不同且无直接因果换算（causalClaimAllowed=false），严禁推论自动停线。",
            "Operations 已登记行动项 A-05（2026-10-16 前完成 LINE-01 影响边界核实），向管理层汇报时将严格区分 70 台总成缺口与 150 件来料缺口。",
        ],
        "【AB-730 20261015 Automotive】跨部门改善行动确认": [
            "依据 05-Action-Tracker.xlsx 与 08-Improvement-Meeting.docx，目前共汇总 A-01 至 A-06 六项行动：A-01、A-02、A-04、A-05、A-06 五项已确认责任角色与到期日；A-03（来源 PO-005，建议角色 Supply）仍为未确认提案（ownerConfirmed=false，dueDate=null，ProposalOnly，D-03）。",
            "Sponsor 批示：请严格保留 A-03 为待确认提案状态，在 Supply 正式确认前不得虚构截止日期。",
            "Planner 收到，已在 05-Action-Tracker.xlsx 与 08-Improvement-Meeting.docx 同步标注 D-03 风险提示。",
        ],
    }

    email_threads = []
    for idx, thread in enumerate(spec["emailThreads"], start=1):
        subject = thread["subject"]
        authors = thread["authors"]
        to_lists = [thread["initialToRoles"], thread["replyToRoles"][0], thread["replyToRoles"][1]]
        bodies = thread_bodies[subject]
        sources_str = ", ".join(thread.get("sources", []))

        msgs = []
        for m_idx in range(3):
            action = "send" if m_idx == 0 else "reply"
            msg_subject = subject if m_idx == 0 else f"Re: {subject}"
            body_html = (
                f"<p><strong>【AB-730 虚构训练资料】</strong>（角色：{authors[m_idx]} | 关联数据：{sources_str}）</p>"
                f"<p>{bodies[m_idx]}</p>"
            )
            msgs.append(
                {
                    "action": action,
                    "fromRole": authors[m_idx],
                    "toRoles": list(to_lists[m_idx]),
                    "ccRoles": list(cc_roles),
                    "subject": msg_subject,
                    "bodyHtml": body_html,
                }
            )
        email_threads.append(
            {
                "threadName": f"T-0{idx} {subject}",
                "emails": msgs,
            }
        )
    return {"emailThreads": email_threads}


def _build_teams_spec(spec: dict) -> dict:
    story = spec["teamStoryTimes"]
    day_offset = int(story["dayOffset"])
    t1 = story["firstPostTime"]
    tr = story["replyTime"]
    t2 = story["secondPostTime"]

    channel_scripts = {
        "供应链交期": (
            "【02-Supplier-Delivery.xlsx / 07-Supplier-Response.docx】五笔采购单计划 1,200 件、实收 1,050 件，交期缺口 150 件；PO-005（P-111，S-02）实收 150 件但实际到货日未核实（actualDate 为空，状态 Unknown，D-01）。",
            "收到。Planner 汇总周报时会将 PO-005 准时状态列为 Unknown，绝不默认归入准时交货。",
            "补充：PO-002 已建立 A-02 跟进（2026-10-16 到期），PO-005 对应的 A-03 仍处于 ProposalOnly（未确认，D-03）。",
        ),
        "BOM与成本": (
            "【01-BOM-Cost.xlsx】虚构局部总成 12 项零件（P-101..P-112）基准成本 2,582 元、现行成本 2,682 元，单台净增 +100 元（CNY，非整车 BOM）。",
            "Analytics 已核对：仅 P-101（用量 2 件）单价由 500 元升至 550 元（+100 元），其余 11 项零件无变化。",
            "Engineering 已关联行动项 A-04（2026-10-18 到期），完成 P-101 价格变动工程签证复核。",
        ),
        "零件质量": (
            "【03-Quality-Inspection.xlsx / 06-Quality-Brief.docx】独立抽检批次 Q-001（1000/20，2.00%）、Q-002（100/10，10.00%）、Q-003（0/0，N/A，D-02）；有效样本 1,100 件、不良 30 件，加权不良率 2.73%（D-04）。",
            "Engineering 确认：Q-003 抽检分母为 0 必须保持 N/A，且总不良率按 30/1100=2.73% 计算，不采用简单算术平均。",
            "Quality 已将 Q-002 异常纳入行动项 A-01（2026-10-17 到期），并由 Analytics 在 A-06 中输出加权口径报告。",
        ),
        "生产影响": (
            "【04-Production-Impact.xlsx】LINE-01（计划 500/实际 450，缺口 50）与 LINE-02（计划 400/实际 380，缺口 20）合计虚构局部总成缺口 70 台。",
            "Quality 会签：70 台局部总成缺口与 150 件采购交期缺口单位不同且无直接因果映射（causalClaimAllowed=false）。",
            "Operations 已登记 A-05（2026-10-16 到期），汇报材料中严禁把 70 台产出缺口直接归因于 150 件交期缺口或自动推论停线。",
        ),
        "改善行动": (
            "【05-Action-Tracker.xlsx / 08-Improvement-Meeting.docx】改善行动共 6 项（A-01..A-06），其中 A-03（来源 PO-005，建议角色 Supply）未确认责任人且无截止日（ProposalOnly，D-03）。",
            "Sponsor 确认：汇报必须区分已确认行动（A-01、A-02、A-04、A-05、A-06）与未确认提案（A-03），不得为 A-03 编造到期日。",
            "Planner 已完成全部五张工作簿与三份 Word 纪要的交叉核对，供管理层审阅。",
        ),
    }

    channels = []
    for ch in spec["teamChannels"]:
        name = ch["name"]
        top1_author, top2_author = ch["topAuthors"]
        reply_author = ch["firstPostReplyAuthor"]
        post1_text, reply_text, post2_text = channel_scripts[name]

        channels.append(
            {
                "channelName": name,
                "channelDescription": f"AB-730 {name}（虚构汽车供应链训练频道）",
                "messages": [
                    {
                        "fromRole": top1_author,
                        "dayOffset": day_offset,
                        "time": t1,
                        "bodyHtml": f"<p><strong>[{top1_author}]</strong> {post1_text}</p>",
                        "mentions": [],
                        "replies": [
                            {
                                "fromRole": reply_author,
                                "dayOffset": day_offset,
                                "time": tr,
                                "bodyHtml": f"<p><strong>[{reply_author}]</strong> {reply_text}</p>",
                                "mentions": [],
                            }
                        ],
                    },
                    {
                        "fromRole": top2_author,
                        "dayOffset": day_offset,
                        "time": t2,
                        "bodyHtml": f"<p><strong>[{top2_author}]</strong> {post2_text}</p>",
                        "mentions": [],
                        "replies": [],
                    },
                ],
            }
        )
    return {"channels": channels}


def _build_sharepoint_spec(spec: dict) -> dict:
    res = spec["resources"]
    member_roles = [rk for rk in spec["roleKeys"] if rk != "Admin"]
    docs = [
        {
            "sourceFilename": fname,
            "targetFilename": fname,
            "libraryName": res["sharePointLibrary"],
        }
        for fname in spec["sharePointDocuments"]
    ]
    return {
        "sites": [
            {
                "alias": res["sharePointAlias"],
                "displayName": res["sharePointDisplayName"],
                "description": "AB-730 Automotive fictional training knowledge source (zh-CN).",
                "visibility": "Private",
                "owners": ["Admin"],
                "members": member_roles,
                "documentLibrary": res["sharePointLibrary"],
                "documents": docs,
                "lists": [],
            }
        ]
    }


def _stabilize_docx_bytes(doc: Document) -> bytes:
    doc.core_properties.author = "AB-730 demo"
    doc.core_properties.comments = "Fictional training data for AB-730 (zh-CN)."
    doc.core_properties.created = datetime(2020, 1, 1)
    doc.core_properties.modified = datetime(2020, 1, 1)

    raw_stream = io.BytesIO()
    doc.save(raw_stream)

    stable_stream = io.BytesIO()
    with ZipFile(io.BytesIO(raw_stream.getvalue()), "r") as src_zf, ZipFile(
        stable_stream, "w"
    ) as dst_zf:
        for name in sorted(src_zf.namelist()):
            info = ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            content = src_zf.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:created[^>]*>)[^<]*",
                    rb"\g<1>2020-01-01T00:00:00Z",
                    content,
                )
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>2020-01-01T00:00:00Z",
                    content,
                )
            dst_zf.writestr(info, content)
    return stable_stream.getvalue()


_DOCX_TABLE_WIDTHS_DXA_BY_COLS: dict[int, list[int]] = {
    6: [1600, 2400, 1500, 1500, 1500, 5900],
    7: [1300, 2300, 4000, 1600, 1500, 1500, 2200],
    9: [1500, 1800, 1600, 1800, 1800, 1400, 1400, 1400, 1700],
}


def _set_xml_attr_element(parent, tag: str, attrs: dict[str, str]):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


def _create_docx_document(
    title: str,
    subtitle: str,
    sections: list[tuple[str, list[str]]],
    table_headers: list[str],
    table_rows: list[list[str]],
) -> bytes:
    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.0)
    normal_style.paragraph_format.space_before = Pt(1)
    normal_style.paragraph_format.space_after = Pt(3)

    # First section: US Letter Portrait (8.5" x 11.0") with bounded margins
    portrait_section = doc.sections[0]
    portrait_section.orientation = WD_ORIENT.PORTRAIT
    portrait_section.page_width = Twips(12240)
    portrait_section.page_height = Twips(15840)
    portrait_section.top_margin = Twips(900)
    portrait_section.bottom_margin = Twips(900)
    portrait_section.left_margin = Twips(1080)
    portrait_section.right_margin = Twips(1080)

    h1 = doc.add_heading(title, level=1)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(4)
    for run in h1.runs:
        run.font.color.rgb = RGBColor(0x17, 0x4A, 0x75)

    disclaimer = doc.add_paragraph(
        "【声明】本文件为 AB-730 课程虚构训练资料（语言：简体中文 zh-CN，"
        "币别：人民币 CNY），仅用于 Microsoft 365 Copilot 教学演示与来源核对，不代表任何真实企业业绩、质量事故或召回结论。"
    )
    if disclaimer.runs:
        disclaimer.runs[0].bold = True

    doc.add_paragraph(subtitle)

    for heading_text, paragraphs in sections:
        h2 = doc.add_heading(heading_text, level=2)
        h2.paragraph_format.space_before = Pt(4)
        h2.paragraph_format.space_after = Pt(2)
        h2.paragraph_format.keep_with_next = True
        for p_text in paragraphs:
            doc.add_paragraph(p_text)

    # Separate Landscape section on a new page before the table: US Letter 11.0" x 8.5" (15840 x 12240 DXA)
    table_section = doc.add_section(WD_SECTION.NEW_PAGE)
    table_section.orientation = WD_ORIENT.LANDSCAPE
    table_section.page_width = Twips(15840)
    table_section.page_height = Twips(12240)
    table_section.top_margin = Twips(720)
    table_section.bottom_margin = Twips(720)
    table_section.left_margin = Twips(720)
    table_section.right_margin = Twips(720)

    table_h2 = doc.add_heading("核心数据核对表", level=2)
    table_h2.paragraph_format.space_before = Pt(0)
    table_h2.paragraph_format.space_after = Pt(4)
    table_h2.paragraph_format.keep_with_next = True

    col_count = len(table_headers)
    col_widths_dxa = _DOCX_TABLE_WIDTHS_DXA_BY_COLS.get(
        col_count, [14400 // max(col_count, 1)] * col_count
    )
    total_table_width_dxa = sum(col_widths_dxa)

    table = doc.add_table(rows=1, cols=col_count)
    table.style = "Table Grid"
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    _set_xml_attr_element(
        tbl_pr, "w:tblW", {"w:w": str(total_table_width_dxa), "w:type": "dxa"}
    )
    _set_xml_attr_element(tbl_pr, "w:tblLayout", {"w:type": "fixed"})

    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for idx, width_dxa in enumerate(col_widths_dxa):
        table.columns[idx].width = Twips(width_dxa)
        if idx < len(grid_cols):
            grid_cols[idx].set(qn("w:w"), str(width_dxa))

    hdr_row = table.rows[0]
    hdr_cells = hdr_row.cells
    for idx, h_text in enumerate(table_headers):
        hdr_cells[idx].text = h_text

    for row_vals in table_rows:
        row_cells = table.add_row().cells
        for idx, val in enumerate(row_vals):
            row_cells[idx].text = str(val)

    for r_idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        _set_xml_attr_element(tr_pr, "w:cantSplit", {})
        if r_idx == 0:
            _set_xml_attr_element(tr_pr, "w:tblHeader", {})

        for c_idx, cell in enumerate(row.cells):
            width_dxa = col_widths_dxa[c_idx]
            cell.width = Twips(width_dxa)
            tc_pr = cell._tc.get_or_add_tcPr()
            _set_xml_attr_element(
                tc_pr, "w:tcW", {"w:w": str(width_dxa), "w:type": "dxa"}
            )
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    if r_idx == 0:
                        run.bold = True

    return _stabilize_docx_bytes(doc)


def _build_three_docx_files(spec: dict) -> dict[str, bytes]:
    inspections = spec["inspections"]
    deliveries = spec["deliveries"]
    bom = spec["bom"]
    production = spec["production"]
    actions = spec["actions"]
    expected = spec["expectedFacts"]

    actions_by_id = {a["id"]: a for a in actions}
    a01 = actions_by_id["A-01"]
    a03 = actions_by_id["A-03"]
    a06 = actions_by_id["A-06"]

    # Build 06-Quality-Brief.docx dynamically from spec["inspections"], spec["expectedFacts"], spec["actions"]
    q_table_rows: list[list[str]] = []
    q_part_ids: list[str] = []
    q_row_notes = {
        "Q-001": "独立生产/库存抽检（有效样本）",
        "Q-002": f'独立生产/库存抽检（触发 {a01["id"]}）',
        "Q-003": f'抽检数为 0 显式返回 {inspections["zeroDenominatorResult"]}（D-02，非 0%）',
    }
    for batch_id, part_id, inspected_qty, defective_qty in inspections["rows"]:
        q_part_ids.append(str(part_id))
        if inspected_qty == 0:
            rate_str = str(inspections["zeroDenominatorResult"])
        else:
            rate_str = f"{(defective_qty / inspected_qty) * 100:.2f}%"
        q_table_rows.append(
            [
                str(batch_id),
                str(part_id),
                str(inspected_qty),
                str(defective_qty),
                rate_str,
                q_row_notes.get(str(batch_id), "独立抽检批次"),
            ]
        )

    valid_inspected = int(expected["validInspectionPieces"])
    total_defective = int(expected["defectivePieces"])
    weighted_rate_str = f'{expected["weightedDefectRatePercentRounded2"]}%'
    q_table_rows.append(
        [
            "加权合计",
            "/".join(q_part_ids),
            str(valid_inspected),
            str(total_defective),
            weighted_rate_str,
            f"{total_defective}/{valid_inspected} 加权口径（D-04，非算术平均）",
        ]
    )

    q_r0, q_r1, q_r2 = inspections["rows"]
    doc_06 = _create_docx_document(
        title="06-Quality-Brief：零件批次质量抽检简报（Q-001 至 Q-003）",
        subtitle="编制角色：Quality / Engineering / Analytics | 关联来源：03-Quality-Inspection.xlsx（inspections, expectedFacts, actions）",
        sections=[
            (
                "一、独立生产/库存抽检范围与加权不良率（D-04）",
                [
                    f'依据 03-Quality-Inspection.xlsx（表格 QualityInspection），抽检范围为独立虚构生产/库存抽检（{inspections["scope"]}），并非 02-Supplier-Delivery.xlsx 五笔采购单（PO-001 至 PO-005）的逐行映射或普查。',
                    f"{q_r0[0]}（零件 {q_r0[1]}）抽检 {q_r0[2]:,} 件、不良 {q_r0[3]} 件（不良率 {(q_r0[3] / q_r0[2]) * 100:.2f}%）；{q_r1[0]}（零件 {q_r1[1]}）抽检 {q_r1[2]:,} 件、不良 {q_r1[3]} 件（不良率 {(q_r1[3] / q_r1[2]) * 100:.2f}%）。",
                    f"有效抽检样本合计 {valid_inspected:,} 件（{valid_inspected}），不良合计 {total_defective} 件，加权不良率为 {total_defective} / {valid_inspected} = {weighted_rate_str}（缺陷标记 D-04，严禁按批次百分比简单算术平均计算）。",
                ],
            ),
            (
                "二、Q-003 零分母规则（D-02）与关联行动项（A-01、A-06）",
                [
                    f'{q_r2[0]}（零件 {q_r2[1]}）属于独立生产/库存待检批次，抽检数量为 {q_r2[2]}、不良数量为 {q_r2[3]}，结果显式返回 {inspections["zeroDenominatorResult"]}（缺陷标记 D-02），不得记为 0% 不良率。',
                    f'关联改善行动：{a01["id"]}（来源 {a01["source"]}，负责角色 {a01["ownerRole"]}，目标完成日 {a01["dueDate"]}，状态 {a01["status"]}）；{a06["id"]}（来源 {a06["source"]}，负责角色 {a06["ownerRole"]}，目标完成日 {a06["dueDate"]}，状态 {a06["status"]}）。',
                ],
            ),
        ],
        table_headers=["抽检批号", "零件编号", "抽检数量", "不良数量", "不良率", "范围与状态说明"],
        table_rows=q_table_rows,
    )

    # Build 07-Supplier-Response.docx dynamically from spec["deliveries"] and spec["expectedFacts"]
    del_table_rows: list[list[str]] = []
    first4_planned = sum(int(r[5]) for r in deliveries["rows"][:4])
    first4_received = sum(int(r[6]) for r in deliveries["rows"][:4])
    gap_po_ids: list[str] = []
    for po_id, part_id, supplier_id, promised_date, actual_date, planned_qty, received_qty in deliveries["rows"]:
        gap_qty = int(planned_qty) - int(received_qty)
        if gap_qty > 0:
            gap_po_ids.append(f"{po_id}（缺口 {gap_qty} 件）")
        if actual_date is None:
            actual_str = "Unknown"
            status_str = "Unknown"
        elif actual_date <= promised_date:
            actual_str = str(actual_date)
            status_str = "OnTime"
        else:
            actual_str = str(actual_date)
            status_str = "Delayed"
        del_table_rows.append(
            [
                str(po_id),
                str(part_id),
                str(supplier_id),
                str(promised_date),
                actual_str,
                str(planned_qty),
                str(received_qty),
                str(gap_qty),
                status_str,
            ]
        )

    total_planned_del = int(expected["plannedDeliveryPieces"])
    total_received_del = int(expected["receivedDeliveryPieces"])
    total_gap_del = int(expected["deliveryGapPieces"])
    po5 = deliveries["rows"][4]
    po5_gap = int(po5[5]) - int(po5[6])
    del_table_rows.append(
        [
            "合计",
            "-",
            "-",
            "-",
            "-",
            str(total_planned_del),
            str(total_received_del),
            str(total_gap_del),
            f"{po5[0]} 到货日未核实（D-01）",
        ]
    )

    doc_07 = _create_docx_document(
        title="07-Supplier-Response：供应商交期与缺口核对回复（PO-001 至 PO-005）",
        subtitle="编制角色：Supply / Operations / Planner | 关联来源：02-Supplier-Delivery.xlsx（deliveries, expectedFacts）",
        sections=[
            (
                "一、五笔采购单交期与缺口分布",
                [
                    f"依据 02-Supplier-Delivery.xlsx（表格 SupplierDelivery），PO-001 至 PO-005 五笔采购单计划交货合计 {total_planned_del:,} 件（{total_planned_del}），实收合计 {total_received_del:,} 件（{total_received_del}），总交期缺口 {total_gap_del} 件（单位：{deliveries['unit']}）。",
                    f"其中前四笔采购单（PO-001 至 PO-004）计划合计 {first4_planned:,} 件（{first4_planned}）、实收合计 {first4_received:,} 件（{first4_received}）；缺口分布于 {'、'.join(gap_po_ids)}。",
                ],
            ),
            (
                "二、PO-005 实际到货日未核实异常说明（D-01）",
                [
                    f"{po5[0]}（零件 {po5[1]}，供应商 {po5[2]}，承诺交期 {po5[3]}）计划数量 {po5[5]} 件、实收数量 {po5[6]} 件、缺口 {po5_gap} 件。",
                    f"由于 {po5[0]} 实际到货日（actualDate）为空值（{deliveries['missingActualDateMeaning']}），准时状态标记为 Unknown（缺陷标记 D-01），严禁在汇总汇报中将 {po5[0]} 默认归入准时到货。",
                ],
            ),
        ],
        table_headers=[
            "采购单号",
            "零件编号",
            "供应商编号",
            "承诺交期",
            "实际到货日",
            "计划数量",
            "实收数量",
            "缺口数量",
            "准时状态",
        ],
        table_rows=del_table_rows,
    )

    # Build 08-Improvement-Meeting.docx dynamically from bom, deliveries, inspections, production, actions
    action_titles = {
        "A-01": "排查 Q-002（P-103）批次不良率 10.00% 根因并复测",
        "A-02": "确认 PO-002（S-02）延迟到货与缺口 50 件补齐安排",
        "A-03": "提案：确认 PO-005（P-111）实际到货日与尾数 50 件追补方案（待确认）",
        "A-04": "核对 P-101 单价由 500 元上调至 550 元（单台 +100 元）技术变更依据",
        "A-05": "核实 LINE-01/LINE-02 合计 70 台排程缺口边界，明确非 150 件交期缺口直接导致",
        "A-06": "按加权口径 30/1100 = 2.73% 输出质量周报（禁止算术平均）",
    }
    act_table_rows: list[list[str]] = []
    for act in actions:
        act_table_rows.append(
            [
                str(act["id"]),
                str(act["source"]),
                action_titles.get(str(act["id"]), str(act["status"])),
                str(act["ownerRole"]),
                "是" if act["ownerConfirmed"] else "否",
                "null" if act["dueDate"] is None else str(act["dueDate"]),
                str(act["status"]),
            ]
        )

    base_cost = int(expected["baseAssemblyCost"])
    curr_cost = int(expected["currentAssemblyCost"])
    delta_cost = int(expected["assemblyCostDelta"])
    prod_gap = int(expected["productionGapAssemblies"])
    line1, line2 = production["rows"]
    prod_plan_total = int(line1[1]) + int(line2[1])
    prod_act_total = int(line1[2]) + int(line2[2])

    doc_08 = _create_docx_document(
        title="08-Improvement-Meeting：跨部门改善模拟会议记录（A-01 至 A-06）",
        subtitle=(
            "文件性质：虚构模拟会议记录 | 参会角色：Sponsor / Planner / Supply / Quality / Operations / Engineering / Analytics | "
            "关联来源：01-BOM-Cost.xlsx、02-Supplier-Delivery.xlsx、03-Quality-Inspection.xlsx、04-Production-Impact.xlsx、05-Action-Tracker.xlsx"
        ),
        sections=[
            (
                "一、跨模块核心数据核对摘要（模拟会议记录）",
                [
                    f"1. BOM 成本（01-BOM-Cost.xlsx）：虚构局部总成 12 项零件（P-101 至 P-112，{bom['unit']}）基准单台成本 {base_cost:,} 元（{base_cost}），现行单台成本 {curr_cost:,} 元（{curr_cost}），净差异 +{delta_cost} 元（仅 P-101 用量 2 件由 500 元升至 550 元）。",
                    f"2. 供应商交期（02-Supplier-Delivery.xlsx）：PO-001 至 PO-005 计划 {total_planned_del:,} 件（{total_planned_del}）、实收 {total_received_del:,} 件（{total_received_del}）、缺口 {total_gap_del} 件（PO-002、PO-004、PO-005 各缺 50 件；{po5[0]} 零件 {po5[1]} 实际到货日 Unknown，D-01）。",
                    f"3. 质量抽检（03-Quality-Inspection.xlsx）：独立抽检批次 Q-001（P-101，1000/20）、Q-002（P-103，100/10）、Q-003（P-108，0/0 返回 N/A，D-02），有效样本 {valid_inspected:,} 件、不良 {total_defective} 件，加权不良率 {weighted_rate_str}（D-04）。",
                    f"4. 产线排程（04-Production-Impact.xlsx）：{line1[0]}（计划 {line1[1]}/实际 {line1[2]}，缺口 {int(line1[1]) - int(line1[2])} 台）与 {line2[0]}（计划 {line2[1]}/实际 {line2[2]}，缺口 {int(line2[1]) - int(line2[2])} 台）合计计划 {prod_plan_total} 台、实际 {prod_act_total} 台，虚构局部总成缺口 {prod_gap} 台（causalClaimAllowed={production['causalClaimAllowed']}，与 {total_gap_del} 件采购交期缺口无直接因果换算，不得虚构停线原因或归因于交期缺口）。",
                ],
            ),
            (
                "二、改善行动追踪与 A-03 未确认提案边界（D-03）",
                [
                    f"依据 05-Action-Tracker.xlsx（表格 ImprovementActions），本模拟会议记录共核对 {expected['actionCount']} 项改善行动（A-01 至 A-06）。",
                    f'其中 {a03["id"]}（来源 {a03["source"]}，建议负责角色 {a03["ownerRole"]}）尚未获责任人确认（ownerConfirmed={a03["ownerConfirmed"]}），目标完成日为空（dueDate=null），状态为 {a03["status"]}（缺陷标记 D-03），汇报时必须单独列为待确认提案，严禁编造截止日期。',
                ],
            ),
        ],
        table_headers=["行动编号", "来源单号", "改善事项", "负责角色", "责任人已确认", "目标完成日", "状态"],
        table_rows=act_table_rows,
    )

    return {
        "06-Quality-Brief.docx": doc_06,
        "07-Supplier-Response.docx": doc_07,
        "08-Improvement-Meeting.docx": doc_08,
    }


def _neutral_config_example(spec: dict) -> dict:
    date = spec["deliveryDate"]
    course = spec["courseCode"]
    return {
        "tenantId": "<tenant-guid>",
        "clientId": "<app-client-id>",
        "clientSecret": "<local-secret>",
        "adminUpn": "admin@example.invalid",
        "demoUserUpn": "admin@example.invalid",
        "teamDisplayName": f"{date}-{course}",
        "teamDescription": f"{course} dated training demo",
        "filesSourceDir": "DEMO-FILE",
        "roles": {
            rk: {
                "upn": f"{rk.lower()}@example.invalid",
                "displayName": rk,
            }
            for rk in spec["roleKeys"]
        },
    }


def _build_expected_results(spec: dict, asset_sha256: dict[str, str]) -> dict:
    return {
        "scenarioSlug": spec["scenarioSlug"],
        "courseCode": spec["courseCode"],
        "deliveryDate": spec["deliveryDate"],
        "datasetAsOf": spec["datasetAsOf"],
        "locale": spec["locale"],
        "currency": spec["currency"],
        "resources": spec["resources"],
        "expectedFacts": spec["expectedFacts"],
        "computedFactIds": spec["computedFactIds"],
        "intentionalAnomalies": spec.get("intentionalAnomalies", []),
        "assetSHA256": asset_sha256,
    }


def _scan_candidate_directory_for_leaks(candidate_dir: Path, forbidden_tokens: list[str]) -> None:
    for path in sorted(candidate_dir.rglob("*")):
        if path.is_dir():
            continue
        if path.suffix in {".json", ".md", ".txt", ".yaml", ".yml", ".ps1"}:
            text = path.read_text(encoding="utf-8")
            for token in forbidden_tokens:
                if token in text:
                    raise ValueError(f"Redaction check failed: forbidden token found in {path.name}")
        elif path.suffix in {".xlsx", ".docx"}:
            with ZipFile(path, "r") as zf:
                for member in zf.namelist():
                    if member.endswith((".xml", ".rels")):
                        xml_text = zf.read(member).decode("utf-8", errors="ignore")
                        for token in forbidden_tokens:
                            if token in xml_text:
                                raise ValueError(
                                    f"Redaction check failed: forbidden token found in {path.name}:{member}"
                                )


def build_ab730_scenario(
    spec_path: Path | str,
    identity_config_path: Path | str,
    date: str,
    output_dir: Path | str,
) -> Path:
    spec_path = Path(spec_path)
    identity_config_path = Path(identity_config_path)
    output_dir = Path(output_dir)

    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    identity_config = json.loads(identity_config_path.read_text(encoding="utf-8"))

    staged_roles, tenant_domain, forbidden_tokens = validate_spec_and_inputs(
        spec, identity_config, str(date), output_dir
    )

    # Assemble synthetic pack in session-private staging (never copying secrets)
    staged_pack = {
        "slug": spec["scenarioSlug"],
        "customer": {
            "name": "AB-730 Fictional Automotive Demo",
            "industry": "Automotive",
            "locale": spec["locale"],
        },
        "course": {
            "code": spec["courseCode"],
            "purpose": "Automotive",
        },
        "tenantDomain": tenant_domain,
        "roles": staged_roles,
        "teams": _build_teams_spec(spec),
        "emails": _build_emails_spec(spec),
        "workbooks": _build_workbooks_spec(spec),
        "sharepoint": _build_sharepoint_spec(spec),
    }

    staging_root = _resolve_staging_root(output_dir)
    run_stage_dir = staging_root / f"stage-{uuid.uuid4().hex}"
    run_stage_dir.mkdir(parents=True, exist_ok=True)
    staged_pack_path = run_stage_dir / "pack.yaml"
    staged_scenario_dir = run_stage_dir / "scenario"

    staged_pack_path.write_text(
        yaml.safe_dump(staged_pack, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )

    builder = _load_build_scenario_module()
    builder.build_scenario(staged_pack_path, str(date), staged_scenario_dir)

    # Generate the 3 canonical .docx files in staging DEMO-FILE
    demo_stage_dir = staged_scenario_dir / "DEMO-FILE"
    for docx_name, docx_bytes in _build_three_docx_files(spec).items():
        (demo_stage_dir / docx_name).write_bytes(docx_bytes)

    # Compute deterministic SHA256 hashes for all 8 Office files in canonical order
    ordered_filenames = [item["name"] for item in spec["officeFiles"]]
    asset_sha256: dict[str, str] = {}
    manifest_files: list[dict[str, str]] = []
    for fname in ordered_filenames:
        fpath = demo_stage_dir / fname
        if not fpath.exists():
            raise FileNotFoundError(f"Expected generated Office asset missing in staging: {fname}")
        digest = hashlib.sha256(fpath.read_bytes()).hexdigest()
        asset_sha256[fname] = digest
        manifest_files.append({"localName": fname, "sha256": digest})

    # Enrich teams-messages.json with top-level teamDisplayName & teamDescription for runner compatibility
    teams_json_path = staged_scenario_dir / "teams-messages.json"
    teams_payload = json.loads(teams_json_path.read_text(encoding="utf-8"))
    teams_enriched = {
        "teamDisplayName": spec["resources"]["teamDisplayName"],
        "teamDescription": f'{spec["courseCode"]} dated training demo',
        "channels": teams_payload["channels"],
    }
    builder.write_json(teams_json_path, teams_enriched)

    # Write updated files-manifest.json, expected-results.json, neutral config.json.example, and README.md
    builder.write_json(
        staged_scenario_dir / "files-manifest.json",
        {
            "targetFolder": spec["resources"]["oneDriveFolder"],
            "uploadToRole": "Admin",
            "files": manifest_files,
            "assetSHA256": asset_sha256,
        },
    )
    builder.write_json(
        staged_scenario_dir / "expected-results.json",
        _build_expected_results(spec, asset_sha256),
    )
    builder.write_json(
        staged_scenario_dir / "config.json.example",
        _neutral_config_example(spec),
    )
    (staged_scenario_dir / "README.md").write_text(
        f'# {spec["courseCode"]} {spec["scenarioSlug"]}\n\n'
        "Fictional demonstration content (`zh-CN`, `CNY`) for AB-730; customer and supplier names "
        "are narrative training fixtures only. Only existing, enabled tenant accounts may be used. "
        "This seeder never creates users or modifies mailboxes destructively.\n\n"
        "Copy `config.json.example` to locally ignored `config.json` (or pass `-ConfigPath`), "
        "then run `pwsh -File .\\run.ps1 -PreflightOnly` before any write phase.\n",
        encoding="utf-8",
    )

    # Fail-closed privacy/redaction check on staged candidate before exporting to output_dir
    _scan_candidate_directory_for_leaks(staged_scenario_dir, forbidden_tokens)

    # Export candidate files to output_dir without touching any existing run.ps1
    output_dir.mkdir(parents=True, exist_ok=True)
    out_demo_dir = output_dir / "DEMO-FILE"
    out_demo_dir.mkdir(parents=True, exist_ok=True)

    for top_name in (
        "config.json.example",
        "emails.json",
        "teams-messages.json",
        "sharepoint-sites.json",
        "files-manifest.json",
        "expected-results.json",
        "README.md",
    ):
        (output_dir / top_name).write_bytes((staged_scenario_dir / top_name).read_bytes())

    for fname in ordered_filenames:
        (out_demo_dir / fname).write_bytes((demo_stage_dir / fname).read_bytes())

    return output_dir


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build AB-730 public scenario assets from canonical spec and staging identity binding."
    )
    parser.add_argument("--spec", required=True, help="Path to ab730-data-spec.json")
    parser.add_argument("--identity-config", required=True, help="Path to identity config/manifest JSON")
    parser.add_argument("--date", required=True, help="Scenario delivery date (YYYYMMDD)")
    parser.add_argument("--output", required=True, help="Target scenario output directory")
    args = parser.parse_args(argv)

    build_ab730_scenario(
        spec_path=args.spec,
        identity_config_path=args.identity_config,
        date=args.date,
        output_dir=args.output,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

