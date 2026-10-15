import datetime as dt
import hashlib
import importlib.util
import json
import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook


GENERATOR_DIR = Path(__file__).parents[1]
AB730_SCRIPT = GENERATOR_DIR / "build_ab730_scenario.py"
SPEC_PATH = Path(__file__).resolve().parent / "ab730-data-spec.json"

ROLE_KEYS = [
    "Admin",
    "Sponsor",
    "Planner",
    "Supply",
    "Quality",
    "Operations",
    "Engineering",
    "Analytics",
]

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def load_ab730_builder():
    if not AB730_SCRIPT.exists():
        raise FileNotFoundError(f"Missing {AB730_SCRIPT}")
    spec = importlib.util.spec_from_file_location("build_ab730_scenario", AB730_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def synthetic_identity_config(domain: str = "example.invalid") -> dict:
    return {
        "tenantId": "00000000-0000-0000-0000-000000000001",
        "clientId": "00000000-0000-0000-0000-000000000002",
        "clientSecret": "super-secret-do-not-copy-12345",
        "adminUpn": f"admin@{domain}",
        "demoUserUpn": f"admin@{domain}",
        "roles": {
            role: {
                "upn": f"{role.lower()}@{domain}",
                "displayName": f"{role} Test Actor",
            }
            for role in ROLE_KEYS
        },
    }


def write_json(path: Path, data: dict) -> Path:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def test_cli_interface_and_end_to_end_generation(tmp_path):
    builder = load_ab730_builder()
    id_cfg_path = write_json(
        tmp_path / "identity-config.json",
        synthetic_identity_config(domain="private-fixture.example.invalid"),
    )
    out_dir = tmp_path / "ab730-out"

    rc = builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )
    assert rc == 0

    expected_top_files = {
        "config.json.example",
        "emails.json",
        "teams-messages.json",
        "sharepoint-sites.json",
        "files-manifest.json",
        "expected-results.json",
        "README.md",
        "DEMO-FILE",
    }
    assert expected_top_files <= {p.name for p in out_dir.iterdir()}

    demo_files = sorted(p.name for p in (out_dir / "DEMO-FILE").iterdir())
    assert demo_files == [
        "01-BOM-Cost.xlsx",
        "02-Supplier-Delivery.xlsx",
        "03-Quality-Inspection.xlsx",
        "04-Production-Impact.xlsx",
        "05-Action-Tracker.xlsx",
        "06-Quality-Brief.docx",
        "07-Supplier-Response.docx",
        "08-Improvement-Meeting.docx",
    ]


def test_pep723_inline_script_metadata_present():
    script_text = AB730_SCRIPT.read_text(encoding="utf-8")
    assert script_text.startswith("# /// script\n")
    for dep in ("pyyaml", "jsonschema", "openpyxl", "python-docx"):
        assert f'"{dep}"' in script_text.splitlines()[1] or f'"{dep}"' in script_text[:250]


def test_output_refuses_any_non_empty_directory_including_run_ps1(tmp_path):
    builder = load_ab730_builder()
    id_cfg_path = write_json(tmp_path / "identity.json", synthetic_identity_config())

    non_empty = tmp_path / "non-empty"
    non_empty.mkdir()
    (non_empty / "existing.txt").write_text("do not overwrite", encoding="utf-8")
    with pytest.raises(FileExistsError):
        builder.main(
            [
                "--spec",
                str(SPEC_PATH),
                "--identity-config",
                str(id_cfg_path),
                "--date",
                "20261015",
                "--output",
                str(non_empty),
            ]
        )

    runner_dir = tmp_path / "with-runner"
    runner_dir.mkdir()
    (runner_dir / "run.ps1").write_text("# pre-existing file\n", encoding="utf-8")
    with pytest.raises(FileExistsError):
        builder.main(
            [
                "--spec",
                str(SPEC_PATH),
                "--identity-config",
                str(id_cfg_path),
                "--date",
                "20261015",
                "--output",
                str(runner_dir),
            ]
        )


@pytest.mark.parametrize(
    "mutate_cfg,bad_date",
    [
        (lambda c: None, "20260929"),
        (lambda c: None, "2026-10-15"),
        (lambda c: c["roles"].pop("Analytics"), "20261015"),
        (
            lambda c: c["roles"]["Sponsor"].update(upn=c["roles"]["Admin"]["upn"]),
            "20261015",
        ),
        (
            lambda c: c["roles"]["Supply"].update(upn="supply@other-domain.example.invalid"),
            "20261015",
        ),
    ],
)
def test_validation_rejects_invalid_date_or_identity_config(tmp_path, mutate_cfg, bad_date):
    builder = load_ab730_builder()
    cfg = synthetic_identity_config()
    mutate_cfg(cfg)
    id_cfg_path = write_json(tmp_path / "bad-identity.json", cfg)
    out_dir = tmp_path / "should-not-exist"

    with pytest.raises(ValueError):
        builder.main(
            [
                "--spec",
                str(SPEC_PATH),
                "--identity-config",
                str(id_cfg_path),
                "--date",
                bad_date,
                "--output",
                str(out_dir),
            ]
        )
    assert not out_dir.exists()


def test_private_binding_never_leaks_to_public_outputs(tmp_path):
    builder = load_ab730_builder()
    private_domain = "secret-tenant.contoso.invalid"
    cfg = synthetic_identity_config(domain=private_domain)
    id_cfg_path = write_json(tmp_path / "private-id.json", cfg)
    out_dir = tmp_path / "public-out"

    builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )

    forbidden_Needles = [
        private_domain,
        "super-secret-do-not-copy-12345",
        "00000000-0000-0000-0000-000000000001",
        "00000000-0000-0000-0000-000000000002",
        "Test Actor",
        "moneyyu.com",
    ]

    for path in out_dir.rglob("*"):
        if path.is_dir():
            continue
        if path.suffix in {".json", ".md", ".ps1"}:
            text = path.read_text(encoding="utf-8")
            for needle in forbidden_Needles:
                assert needle not in text, f"Forbidden token {needle!r} leaked in {path.name}"
        elif path.suffix in {".xlsx", ".docx"}:
            with ZipFile(path, "r") as zf:
                for member in zf.namelist():
                    if member.endswith(".xml") or member.endswith(".rels"):
                        xml_text = zf.read(member).decode("utf-8", errors="ignore")
                        for needle in forbidden_Needles:
                            assert (
                                needle not in xml_text
                            ), f"Forbidden token {needle!r} leaked in {path.name}:{member}"

    example_cfg = json.loads((out_dir / "config.json.example").read_text(encoding="utf-8"))
    assert example_cfg["tenantId"] == "<tenant-guid>"
    assert example_cfg["clientId"] == "<app-client-id>"
    assert example_cfg["clientSecret"] == "<local-secret>"
    assert example_cfg["adminUpn"] == "admin@example.invalid"
    assert example_cfg["teamDisplayName"] == "20261015-AB-730"
    assert set(example_cfg["roles"].keys()) == set(ROLE_KEYS)
    for role, info in example_cfg["roles"].items():
        assert info["upn"].endswith("@example.invalid")
        assert info["displayName"] == role


def test_staging_root_never_writes_upns_into_output_ancestor_git_tree_and_scans_run_ps1(
    tmp_path, monkeypatch
):
    builder = load_ab730_builder()
    private_domain = "secret-git-boundary.contoso.invalid"
    private_dir = tmp_path / "private-inputs"
    private_dir.mkdir()
    id_cfg_path = write_json(
        private_dir / "private-id.json",
        synthetic_identity_config(domain=private_domain),
    )

    repo_root = tmp_path / "public-git-repo"
    (repo_root / ".git").mkdir(parents=True)
    out_dir = repo_root / "seed-data" / "scenarios" / "ab730-automotive-20261015"

    # Simulate standalone clone where session files/ does not exist
    monkeypatch.setattr(builder, "HERE", repo_root / "seed-data" / "generator")

    resolved_staging = builder._resolve_staging_root(out_dir)
    assert repo_root.resolve() not in [resolved_staging.resolve(), *resolved_staging.resolve().parents]
    assert out_dir.resolve().parent != resolved_staging.resolve().parent

    builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )

    # Walk the entire ancestor Git working tree to prove zero UPNs or private staging leaked anywhere in repo_root
    forbidden_tokens = [
        private_domain,
        f"admin@{private_domain}",
        "super-secret-do-not-copy-12345",
        "Test Actor",
    ]
    for path in repo_root.rglob("*"):
        if path.is_dir() or ".git" in path.parts:
            continue
        if path.suffix in {".json", ".md", ".txt", ".yaml", ".yml", ".ps1"}:
            text = path.read_text(encoding="utf-8")
            for token in forbidden_tokens:
                assert token not in text, f"Forbidden token {token!r} found in Git tree file {path}"

    # Verify fail-closed if no safe staging directory outside Git tree can be resolved
    with pytest.raises(ValueError):
        builder._assert_safe_private_staging_root(
            out_dir.parent / ".ab730-generator-staging", out_dir
        )

    # Verify _scan_candidate_directory_for_leaks scans run.ps1 and does not skip it
    leak_dir = tmp_path / "leak-check-dir"
    leak_dir.mkdir()
    (leak_dir / "run.ps1").write_text(
        f"$admin = 'admin@{private_domain}'\n", encoding="utf-8"
    )
    with pytest.raises(ValueError, match="run.ps1"):
        builder._scan_candidate_directory_for_leaks(leak_dir, [private_domain])



def test_five_xlsx_workbooks_canonical_math_dates_formats_and_anomalies(tmp_path):
    builder = load_ab730_builder()
    id_cfg_path = write_json(tmp_path / "identity.json", synthetic_identity_config())
    out_dir = tmp_path / "out"
    builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )
    demo_dir = out_dir / "DEMO-FILE"

    # Check all 5 workbooks for creator, named tables, no formula errors
    expected_tables = {
        "01-BOM-Cost.xlsx": ("BOMCost", "A1:G13"),
        "02-Supplier-Delivery.xlsx": ("SupplierDelivery", "A1:J6"),
        "03-Quality-Inspection.xlsx": ("QualityInspection", "A1:F4"),
        "04-Production-Impact.xlsx": ("ProductionImpact", "A1:F3"),
        "05-Action-Tracker.xlsx": ("ImprovementActions", "A1:G7"),
    }
    for fname, (tbl_name, tbl_ref) in expected_tables.items():
        wb = load_workbook(demo_dir / fname, data_only=False)
        assert wb.properties.creator == "AB-730 demo"
        ws = wb.worksheets[0]
        assert tbl_name in ws.tables
        assert ws.tables[tbl_name].ref == tbl_ref
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if isinstance(val, str):
                    for err in ERROR_TOKENS:
                        assert err not in val, f"Error token {err} in {fname}"

    # 01-BOM-Cost.xlsx: 12 rows, base=2582, current=2682, delta=100
    wb_bom = load_workbook(demo_dir / "01-BOM-Cost.xlsx", data_only=False)
    ws_bom = wb_bom.worksheets[0]
    base_sum = sum(ws_bom.cell(r, 2).value * ws_bom.cell(r, 3).value for r in range(2, 14))
    curr_sum = sum(ws_bom.cell(r, 2).value * ws_bom.cell(r, 4).value for r in range(2, 14))
    assert base_sum == 2582
    assert curr_sum == 2682
    assert curr_sum - base_sum == 100
    assert ws_bom["E14"].value == "=SUM(E2:E13)"
    assert ws_bom["F14"].value == "=SUM(F2:F13)"
    assert ws_bom["G14"].value == "=SUM(G2:G13)"

    # 02-Supplier-Delivery.xlsx: real Excel dates, PO-005 actualDate None + Unknown, 1200/1050/150, 0.00%
    wb_del = load_workbook(demo_dir / "02-Supplier-Delivery.xlsx", data_only=False)
    ws_del = wb_del.worksheets[0]
    for r in range(2, 7):
        assert isinstance(ws_del.cell(r, 4).value, dt.date)
        assert ws_del.cell(r, 4).number_format == "yyyy-mm-dd"
        assert ws_del.cell(r, 10).number_format == "0.00%"
    for r in range(2, 6):
        assert isinstance(ws_del.cell(r, 5).value, dt.date)
        assert ws_del.cell(r, 5).number_format == "yyyy-mm-dd"
    assert ws_del["A6"].value == "PO-005"
    assert ws_del["E6"].value is None
    assert ws_del["I6"].value == "Unknown"
    assert sum(ws_del.cell(r, 6).value for r in range(2, 7)) == 1200
    assert sum(ws_del.cell(r, 7).value for r in range(2, 7)) == 1050

    # 03-Quality-Inspection.xlsx: Q-001 1000/20, Q-002 100/10, Q-003 0/0 -> N/A, weighted 30/1100 = 2.73%
    wb_q = load_workbook(demo_dir / "03-Quality-Inspection.xlsx", data_only=False)
    ws_q = wb_q.worksheets[0]
    assert (ws_q["C2"].value, ws_q["D2"].value) == (1000, 20)
    assert (ws_q["C3"].value, ws_q["D3"].value) == (100, 10)
    assert (ws_q["C4"].value, ws_q["D4"].value) == (0, 0)
    assert ws_q["E4"].value == '=IF(C4=0,"N/A",D4/C4)'
    assert ws_q["E5"].value == '=IF(SUM(C2:C4)=0,"N/A",SUM(D2:D4)/SUM(C2:C4))'
    assert ws_q["E2"].number_format == "0.00%"
    assert ws_q["E5"].number_format == "0.00%"
    assert round((20 + 10) / (1000 + 100), 4) == 0.0273

    # 04-Production-Impact.xlsx: 2 lines, 500/450 + 400/380 -> gap 70, not causally mapped to 150
    wb_prod = load_workbook(demo_dir / "04-Production-Impact.xlsx", data_only=False)
    ws_prod = wb_prod.worksheets[0]
    assert sum(ws_prod.cell(r, 2).value for r in range(2, 4)) == 900
    assert sum(ws_prod.cell(r, 3).value for r in range(2, 4)) == 830
    assert ws_prod["E2"].number_format == "0.00%"
    assert ws_prod["E3"].number_format == "0.00%"
    notes_text = " ".join(
        str(ws_prod.cell(r, 6).value or "") for r in range(2, 4)
    ) + str(wb_prod["演示说明"]["A2"].value or "")
    assert "150" in notes_text

    # 05-Action-Tracker.xlsx: 6 rows A-01..A-06, A-03 unconfirmed / null dueDate / ProposalOnly
    wb_act = load_workbook(demo_dir / "05-Action-Tracker.xlsx", data_only=False)
    ws_act = wb_act.worksheets[0]
    assert ws_act["A4"].value == "A-03"
    assert ws_act["B4"].value == "PO-005"
    assert ws_act["D4"].value == "Supply"
    assert ws_act["E4"].value in (False, "否", "Unconfirmed")
    assert ws_act["F4"].value is None
    assert ws_act["G4"].value == "ProposalOnly"
    for r in (2, 3, 5, 6, 7):
        assert isinstance(ws_act.cell(r, 6).value, dt.date)
        assert ws_act.cell(r, 6).number_format == "yyyy-mm-dd"


def test_three_docx_documents_structure_and_disclaimers(tmp_path):
    builder = load_ab730_builder()
    spec = json.loads(SPEC_PATH.read_text(encoding="utf-8"))

    # Verify _build_three_docx_files accepts spec and derives content dynamically from spec
    docx_map = builder._build_three_docx_files(spec)
    assert set(docx_map.keys()) == {
        "06-Quality-Brief.docx",
        "07-Supplier-Response.docx",
        "08-Improvement-Meeting.docx",
    }

    id_cfg_path = write_json(tmp_path / "identity.json", synthetic_identity_config())
    out_dir = tmp_path / "out"
    builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )
    demo_dir = out_dir / "DEMO-FILE"

    forbidden_invented = (
        "中国大陆",
        "多工序换模",
        "工装调校",
        "工装夹具返修",
        "激光雷达支架",
        "感知模块",
        "华东精密传感器",
        "远景汽车制造",
        "电驱总成线",
        "智能座舱线",
        "Q-003 对应 PO-005",
        "900 件、实收 900 件",
        "唯一缺口",
    )

    for fname in ("06-Quality-Brief.docx", "07-Supplier-Response.docx", "08-Improvement-Meeting.docx"):
        doc = Document(demo_dir / fname)
        assert doc.core_properties.author == "AB-730 demo"
        assert len(doc.paragraphs) >= 5
        assert len(doc.tables) == 1
        full_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
            cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells
        )
        assert "虚构" in full_text
        for bad in forbidden_invented:
            assert bad not in full_text, f"Forbidden invented phrase {bad!r} found in {fname}"

    # 1. Verify 06-Quality-Brief.docx table rows match canonical inspections vector
    doc06 = Document(demo_dir / "06-Quality-Brief.docx")
    text06 = "\n".join(p.text for p in doc06.paragraphs)
    assert "独立" in text06 and "非" in text06
    assert "A-01" in text06 and "PendingInvestigation" in text06
    assert "A-06" in text06 and "PendingReport" in text06
    rows06 = [[c.text.strip() for c in r.cells] for r in doc06.tables[0].rows]
    assert [r[:5] for r in rows06[1:4]] == [
        ["Q-001", "P-101", "1000", "20", "2.00%"],
        ["Q-002", "P-103", "100", "10", "10.00%"],
        ["Q-003", "P-108", "0", "0", "N/A"],
    ]
    assert rows06[4][:5] == ["加权合计", "P-101/P-103/P-108", "1100", "30", "2.73%"]

    # 2. Verify 07-Supplier-Response.docx table rows match canonical deliveries vector
    doc07 = Document(demo_dir / "07-Supplier-Response.docx")
    text07 = "\n".join(p.text for p in doc07.paragraphs)
    assert "1,000" in text07 or "1000" in text07
    assert "900" in text07
    assert "PO-002" in text07 and "PO-004" in text07 and "PO-005" in text07
    assert "P-111" in text07 and "2026-10-14" in text07 and "Unknown" in text07
    rows07 = [[c.text.strip() for c in r.cells] for r in doc07.tables[0].rows]
    assert rows07[1:6] == [
        ["PO-001", "P-101", "S-01", "2026-10-10", "2026-10-10", "300", "300", "0", "OnTime"],
        ["PO-002", "P-103", "S-02", "2026-10-11", "2026-10-13", "300", "250", "50", "Delayed"],
        ["PO-003", "P-104", "S-03", "2026-10-12", "2026-10-12", "200", "200", "0", "OnTime"],
        ["PO-004", "P-108", "S-04", "2026-10-13", "2026-10-14", "200", "150", "50", "Delayed"],
        ["PO-005", "P-111", "S-02", "2026-10-14", "Unknown", "200", "150", "50", "Unknown"],
    ]
    assert rows07[6][0] == "合计"
    assert rows07[6][5:8] == ["1200", "1050", "150"]

    # 3. Verify 08-Improvement-Meeting.docx explicit 模拟会议记录 and canonical actions table rows
    doc08 = Document(demo_dir / "08-Improvement-Meeting.docx")
    text08 = "\n".join(p.text for p in doc08.paragraphs)
    assert "模拟会议记录" in text08
    assert "2,582" in text08 and "2,682" in text08 and "100" in text08
    assert "1,200" in text08 and "1,050" in text08 and "150" in text08
    assert "2.73%" in text08 and "70" in text08
    rows08 = [[c.text.strip() for c in r.cells] for r in doc08.tables[0].rows]
    actual_actions_vector = [
        [r[0], r[1], r[3], r[4], r[5], r[6]] for r in rows08[1:]
    ]
    assert actual_actions_vector == [
        ["A-01", "Q-002", "Quality", "是", "2026-10-17", "PendingInvestigation"],
        ["A-02", "PO-002", "Supply", "是", "2026-10-16", "PendingSupplierConfirmation"],
        ["A-03", "PO-005", "Supply", "否", "null", "ProposalOnly"],
        ["A-04", "P-101", "Engineering", "是", "2026-10-18", "PendingCostVerification"],
        ["A-05", "LINE-01", "Operations", "是", "2026-10-16", "PendingImpactVerification"],
        ["A-06", "QualityWeightedRate", "Analytics", "是", "2026-10-16", "PendingReport"],
    ]


def test_three_docx_landscape_table_layout_dxa_widths_cantsplit_and_header_repeat(tmp_path):
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn

    builder = load_ab730_builder()
    id_cfg_path = write_json(tmp_path / "identity.json", synthetic_identity_config())
    out_dir = tmp_path / "out"
    builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )
    demo_dir = out_dir / "DEMO-FILE"

    expected_cols = {
        "06-Quality-Brief.docx": 6,
        "07-Supplier-Response.docx": 9,
        "08-Improvement-Meeting.docx": 7,
    }

    for fname, col_count in expected_cols.items():
        doc = Document(demo_dir / fname)
        assert len(doc.sections) >= 2, f"{fname} must have a separate landscape section for the table"
        table_section = doc.sections[-1]
        assert table_section.orientation == WD_ORIENT.LANDSCAPE
        assert table_section.page_width.twips == 15840
        assert table_section.page_height.twips == 12240
        printable_width_dxa = (
            table_section.page_width.twips
            - table_section.left_margin.twips
            - table_section.right_margin.twips
        )
        assert printable_width_dxa >= 14000

        table = doc.tables[0]
        assert table.autofit is False
        tbl_pr = table._tbl.tblPr
        tbl_layout = tbl_pr.find(qn("w:tblLayout"))
        assert tbl_layout is not None and tbl_layout.get(qn("w:type")) == "fixed"

        grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
        assert len(grid_cols) == col_count
        grid_widths = [int(gc.get(qn("w:w"))) for gc in grid_cols]
        assert sum(grid_widths) <= printable_width_dxa
        assert all(w >= 1200 for w in grid_widths), f"Column widths too narrow in {fname}: {grid_widths}"

        # Verify header row repeats and every row has cantSplit and consistent cell DXA widths
        for r_idx, row in enumerate(table.rows):
            tr_pr = row._tr.find(qn("w:trPr"))
            assert tr_pr is not None, f"Missing trPr on row {r_idx} in {fname}"
            assert tr_pr.find(qn("w:cantSplit")) is not None, f"Missing cantSplit on row {r_idx} in {fname}"
            if r_idx == 0:
                assert tr_pr.find(qn("w:tblHeader")) is not None, f"Missing tblHeader on header row in {fname}"

            for c_idx, cell in enumerate(row.cells):
                tc_pr = cell._tc.find(qn("w:tcPr"))
                assert tc_pr is not None
                tc_w = tc_pr.find(qn("w:tcW"))
                assert tc_w is not None
                assert tc_w.get(qn("w:type")) == "dxa"
                assert int(tc_w.get(qn("w:w"))) == grid_widths[c_idx]



def test_emails_teams_sharepoint_and_manifests_contract(tmp_path):
    builder = load_ab730_builder()
    id_cfg_path = write_json(tmp_path / "identity.json", synthetic_identity_config())
    out_dir = tmp_path / "out"
    builder.main(
        [
            "--spec",
            str(SPEC_PATH),
            "--identity-config",
            str(id_cfg_path),
            "--date",
            "20261015",
            "--output",
            str(out_dir),
        ]
    )

    spec = json.loads(SPEC_PATH.read_text(encoding="utf-8"))

    # Emails: 5 threads x 3 messages, Admin CC'd on all 15
    emails_data = json.loads((out_dir / "emails.json").read_text(encoding="utf-8"))
    assert len(emails_data["emailThreads"]) == 5
    for idx, thread in enumerate(emails_data["emailThreads"]):
        spec_thread = spec["emailThreads"][idx]
        assert len(thread["emails"]) == 3
        assert thread["emails"][0]["subject"] == spec_thread["subject"]
        expected_to = [
            spec_thread["initialToRoles"],
            spec_thread["replyToRoles"][0],
            spec_thread["replyToRoles"][1],
        ]
        for m_idx, msg in enumerate(thread["emails"]):
            assert msg["fromRole"] == spec_thread["authors"][m_idx]
            assert msg["toRoles"] == expected_to[m_idx]
            assert "Admin" in msg["ccRoles"]

    # Teams: 5 channels x (2 posts + 1 reply = 3 messages), date 2026-10-14 (+08:00)
    teams_data = json.loads((out_dir / "teams-messages.json").read_text(encoding="utf-8"))
    assert len(teams_data["channels"]) == 5
    for idx, ch in enumerate(teams_data["channels"]):
        spec_ch = spec["teamChannels"][idx]
        assert ch["channelName"] == spec_ch["name"]
        assert len(ch["messages"]) == 2
        assert len(ch["messages"][0]["replies"]) == 1
        assert ch["messages"][0]["fromRole"] == spec_ch["topAuthors"][0]
        assert ch["messages"][0]["replies"][0]["fromRole"] == spec_ch["firstPostReplyAuthor"]
        assert ch["messages"][1]["fromRole"] == spec_ch["topAuthors"][1]
        assert ch["messages"][0]["createdDateTime"] == "2026-10-14T10:00:00+08:00"
        assert ch["messages"][0]["replies"][0]["createdDateTime"] == "2026-10-14T10:15:00+08:00"
        assert ch["messages"][1]["createdDateTime"] == "2026-10-14T11:00:00+08:00"

    # SharePoint: 1 private site, DemoKnowledge, 3 docx documents, 0 lists
    sp_data = json.loads((out_dir / "sharepoint-sites.json").read_text(encoding="utf-8"))
    assert len(sp_data["sites"]) == 1
    site = sp_data["sites"][0]
    assert site["alias"] == "ab730-automotive-20261015"
    assert site["displayName"] == "AB-730 - Automotive Knowledge - 20261015"
    assert site["visibility"] == "Private"
    assert site["owners"] == ["Admin"]
    assert site["documentLibrary"] == "DemoKnowledge"
    assert site["lists"] == []
    assert [d["sourceFilename"] for d in site["documents"]] == [
        "06-Quality-Brief.docx",
        "07-Supplier-Response.docx",
        "08-Improvement-Meeting.docx",
    ]

    # files-manifest.json and expected-results.json SHA256 verification
    manifest = json.loads((out_dir / "files-manifest.json").read_text(encoding="utf-8"))
    expected_res = json.loads((out_dir / "expected-results.json").read_text(encoding="utf-8"))
    assert manifest["targetFolder"] == "AB730-Automotive-20261015"
    assert manifest["uploadToRole"] == "Admin"
    assert len(manifest["files"]) == 8
    assert "assetSHA256" in manifest
    assert "assetSHA256" in expected_res

    for entry in manifest["files"]:
        name = entry["localName"]
        file_bytes = (out_dir / "DEMO-FILE" / name).read_bytes()
        digest = hashlib.sha256(file_bytes).hexdigest()
        assert entry["sha256"].lower() == digest
        assert manifest["assetSHA256"][name].lower() == digest
        assert expected_res["assetSHA256"][name].lower() == digest


def test_deterministic_asset_hashes_across_runs(tmp_path):
    builder = load_ab730_builder()
    id_cfg_path = write_json(tmp_path / "identity.json", synthetic_identity_config())
    out_a = tmp_path / "run-a"
    out_b = tmp_path / "run-b"

    builder.main(["--spec", str(SPEC_PATH), "--identity-config", str(id_cfg_path), "--date", "20261015", "--output", str(out_a)])
    builder.main(["--spec", str(SPEC_PATH), "--identity-config", str(id_cfg_path), "--date", "20261015", "--output", str(out_b)])

    for p_a in sorted(out_a.rglob("*")):
        if p_a.is_dir():
            continue
        rel = p_a.relative_to(out_a)
        p_b = out_b / rel
        assert p_a.read_bytes() == p_b.read_bytes(), f"Non-deterministic output in {rel}"
