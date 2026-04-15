# -*- coding: utf-8 -*-
"""
AB-730 英業達課程 — 建立 QBR 預演會議逐字稿
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))

def set_cell_shading(cell, color):
    shading_elm = cell._tc.get_or_add_tcPr()
    shading = shading_elm.makeelement(qn('w:shd'), {
        qn('w:fill'): color, qn('w:val'): 'clear',
    })
    shading_elm.append(shading)

def add_styled_table(doc, headers, rows, header_color="2F5496"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(cell, header_color)
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
    return table

def create_qbr_dryrun_transcript():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft JhengHei'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')

    title = doc.add_heading('', level=0)
    run = title.add_run('英業達股份有限公司')
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(47, 84, 150)

    sub = doc.add_heading('Dell QBR 2025 內部預演會議（逐字稿）', level=1)
    for run in sub.runs:
        run.font.color.rgb = RGBColor(47, 84, 150)

    info_lines = [
        ("日期", "2026年1月5日（星期一）14:00-15:30"),
        ("地點", "桃園總部 5F 會議室 A / Microsoft Teams 線上同步"),
        ("會議類型", "QBR 內部預演 (Dry Run)"),
        ("會議目的", "預演 1/15 Dell QBR 簡報，確認內容完整性，練習 Q&A 應對"),
        ("主席", "James Chen (資深專案經理)"),
        ("正式 QBR 日期", "2026/1/15，Dell VP Michael Thompson 主持"),
    ]
    for label, val in info_lines:
        p = doc.add_paragraph()
        run1 = p.add_run(f"{label}：")
        run1.bold = True
        run2 = p.add_run(val)

    doc.add_paragraph()
    doc.add_heading('出席者', level=2)
    attendees = [
        ("James Chen", "PM", "專案管理部", "資深專案經理（QBR 主講）"),
        ("Lisa Wang", "SCM", "供應鏈管理部", "資深供應鏈經理"),
        ("Amy Hsu", "QA", "品質保證部", "品質工程經理"),
        ("Tom Lin", "IT", "資訊系統部", "IT 系統工程師"),
        ("Sarah Chang", "BPM", "流程管理部", "流程改善專員"),
        ("Robert Tsai", "Mfg", "製造部", "製造副總（模擬客戶角色）"),
    ]
    add_styled_table(doc, ["姓名", "部門代碼", "部門", "職稱/角色"], attendees)

    doc.add_paragraph()
    doc.add_heading('會議逐字紀錄', level=2)

    transcript = [
        ("14:02", "James Chen", """各位下午好，謝謝大家。今天是 Dell QBR 的 dry run，1/15 就是正式上場。今天的目標有三個：第一，完整走一遍簡報；第二，找出簡報中的弱點和數據缺口；第三，預演客戶可能問的刁鑽問題。

特別感謝 Robert 副總今天扮演 Dell VP Michael Thompson 的角色，等一下簡報結束後請他用客戶的角度來質問我們。

先快速看一下議程：
1. 簡報走演 (40 min)
2. Robert 副總模擬客戶 Q&A (20 min)
3. 修改討論 (20 min)
4. 行動項目確認 (10 min)

好，那我開始。"""),

        ("14:05", "James Chen", """第一頁，封面。Inventec x Dell — 2025 Annual Business Review。

第二頁，Executive Summary。全年出貨 294,400 台，YoY +12%。營收 USD 485M。三個產品線都有成長，IoT Gateway 成長最快 +25%。

但是——接下來是壞消息——Q3-Q4 的品質和交期明顯下滑。ProServer X100 的交期從 Q2 的 97.5% 掉到 Q4 的 91.3%，DPPM 從 150 升到 280。這是 Michael 最關注的部分，我預計他會在這頁就開始追問。"""),

        ("14:08", "Lisa Wang", """James，插一句。Executive Summary 那頁，建議把「壞消息」放在「好消息」後面，但不要用對比的方式呈現。改用「挑戰與改善」的框架。比如：「Despite strong growth, we faced operational challenges in Q3-Q4 that we have already addressed.」這樣語氣比較主動。"""),

        ("14:09", "James Chen", """好建議，我改一下措辭。

第三到五頁，各產品線績效。ProServer X100 出貨 55,500 台，CloudBook L15 出貨 200,000 台，IoT Gateway 出貨 38,900 台。每個產品線都有交期和品質的季度趨勢圖。

問題是 ProServer X100 那頁——圖表顯示交期從 Q2 開始持續下滑，到 Q4 低於 95% 的紅線。Michael 一定會問為什麼。"""),

        ("14:12", "Amy Hsu", """品質 Deep Dive 那頁我來報告。DPPM 趨勢圖我已經做好了——Q1 180、Q2 150、Q3 220、Q4 280。ProServer X100 惡化最明顯。

我準備了 root cause 的 Pareto chart：
- GPU 散熱片組裝偏移佔 35%——已經導入定位治具，12 月新產出不良率從 4.5% 降到 0.3%
- BMC 韌體問題佔 25%——v2.3.2 已釋出，Dell 端驗證中
- 新進作業員訓練不足佔 20%——11 月已重新排訓
- DRAM 來料問題佔 12%——MemoryPlus 已加強出貨檢驗

9 件客訴的 8D 狀態：4 件已關閉、3 件進行中（1月底前關閉）、1 件待處理（NIC 問題等 Broadcom）、1 件新增（IoT Wi-Fi）。

最擔心的是 QA-2025-055 NIC 相容性問題，Broadcom 那邊回應很慢。如果 Michael 問到這件，我建議的回答是：「We are working closely with Broadcom. The firmware fix is targeted for mid-January. In the interim, we have implemented a workaround by qualifying an alternative NIC from Intel.」"""),

        ("14:18", "Lisa Wang", """交期 Deep Dive。月度交期趨勢圖顯示 10 月是最低點 88.5%。

三大根因：
1. 10 月 MemoryPlus DRAM 來料不良（8.5% reject rate），影響 380 台，這是最大的 hit
2. 11 月 Dell 臨時追加 15% 訂單，我們產能吃不下
3. Q4 NovaTech GPU 交期從 8 週拉到 10 週

改善措施：
- DRAM 安全庫存從 2 週提到 4 週——12 月已生效
- NovaTech Q1 產能保證合約已簽
- 建立 SiliconEdge 備援——Q1 完成驗證
- 計畫增加一條組裝線——投資 $1.2M，Q2 可用

James，我建議在這頁加一個 before/after 的比較圖——左邊是 Q4 的問題，右邊是我們實施的改善和預期效果。這比純文字列表更有說服力。"""),

        ("14:23", "Tom Lin", """數據方面我補充一下。成本節省那頁，全年實績 USD 1,050,000，目標 1,200,000，達成率 87.5%。差距主要來自 Q3-Q4 品質異常的 rework 和 scrap 成本——大約吃掉 180,000 的預計節省。

建議：不要只報一個 87.5% 的數字，要拆開來說——如果扣除品質異常的額外成本，我們的「正常」成本節省其實達到了 102.5%。這樣 Michael 才知道問題出在品質，不是出在我們不努力做 cost down。"""),

        ("14:26", "Sarah Chang", """流程改善那頁，四大成果我都準備好了，每個都有量化數據。

但我想提一個風險——2026 改善提案的投資金額合計 $1.2M（AI 品檢 $600K + SPC $200K + Digital Twin $400K），這個數字 Michael 看到可能會問「你們要花我們多少錢？」。

建議在簡報中明確說明這是英業達自己的投資，不會轉嫁到 Dell 的報價中。而且要算出 ROI——我的估算是 3 年回收 $4.5M，ROI 275%。這個數字非常漂亮。"""),

        ("14:30", "James Chen", """好，謝謝大家的回報。簡報走演完成。現在請 Robert 副總扮演 Dell VP，用最嚴格的角度來質問我們。"""),

        ("14:31", "Robert Tsai", """OK，我現在是 Dell 的 Michael Thompson。James，我看完你的簡報有以下問題：

第一，你說 Q4 交期只有 91.3%，但我看你 10 月只有 88.5%。一個月跌掉快 10 個百分點，你能告訴我那個月到底發生了什麼事嗎？你們的內控出了什麼問題？"""),

        ("14:32", "James Chen", """Michael，非常好的問題。10 月的下跌是由一個特定事件驅動的——我們的 DRAM 供應商有一批來料品質問題，不良率高達 8.5%，遠高於正常的 1% 以下。這批料影響了約 380 台的組裝時程。

我們已經採取了三個措施：第一，DRAM 安全庫存翻倍到 4 週；第二，要求供應商加強出貨檢驗並提供 COC；第三，我們正在建立備援 DRAM 供應商。從 11 月開始，您可以看到交期已經在回升——91.0%、93.8%。"""),

        ("14:34", "Robert Tsai", """第二個問題。你的競爭對手 Foxconn 上季交期是 98.2%，Quanta 是 97.5%。你們 91.3%，這個差距太大了。你怎麼向我保證 Q1 能回到 97% 以上？"""),

        ("14:35", "James Chen", """理解您的 concern。91.3% 確實不是我們的正常水準——前兩季我們都在 96-97%。Q3-Q4 是多個問題同時發生的結果。

關於 Q1 的改善承諾，我有具體的數字支撐：
- DRAM 安全庫存已增加到 4 週——12 月已生效
- GPU 供應商 Q1 產能保證合約已簽——交期回到 8 週
- GPU 散熱片組裝治具已導入——12 月不良率從 4.5% 降到 0.3%
- BMC 韌體 v2.3.2 已釋出並驗證中

基於這些措施，我們承諾 Q1 2026 交期達成率 > 97%，DPPM < 150。如果未達標，我們願意每月提供 progress update 直到恢復。"""),

        ("14:37", "Robert Tsai", """第三個問題。QA-2025-055 的 NIC 問題，你說等 Broadcom 的 fix。如果 Broadcom 一月中給不了，你的 Plan B 是什麼？"""),

        ("14:38", "Amy Hsu", """Michael，這件我來回答。Plan B 我們已經在準備了——我們已經開始驗證 Intel 的 25GbE NIC 作為替代方案。Intel NIC 的驗證預計兩週內完成。如果 Broadcom 在 1 月 15 號前無法提供 firmware fix，我們會切換到 Intel NIC。

另外補充一點——受影響的 310 台我們已經做了隔離，不會出貨給您。維修備品我們也已經準備好了。"""),

        ("14:40", "Robert Tsai", """最後一個問題。你說明年要投資 $1.2M 做 AI 品檢和 Digital Twin。但你今年的表現讓我質疑你們的執行力。你怎麼讓我相信你明年真的會做到？"""),

        ("14:41", "James Chen", """非常公平的問題。我想用今年的實績來回答。今年我們完成了四個流程改善專案——AOI 自動檢測減少誤判 35%、SMT 換線時間縮短 20%、供應商 Dashboard 讓反應時間從 48 小時縮短到 12 小時、AI 預測性維護減少停機 15%。這些都有量化的數據驗證。

明年的投資計畫，我們會設定季度里程碑，每季 QBR 向您報告進度。如果任何專案延遲，我們會立即通知並提供替代方案。

Sarah，可以補充一下 ROI 的部分嗎？"""),

        ("14:43", "Sarah Chang", """好的。三個專案的合計投資 $1.2M，預估 3 年回收 $4.5M，ROI 275%。其中 AI 品檢預計 DPPM 降低 40%，這直接減少貴司的 incoming defect 和維修成本。所有投資由英業達承擔，不會轉嫁到報價中。"""),

        ("14:45", "Robert Tsai", """好，模擬質詢到這裡。脫下 Michael 的帽子，回到 Robert 副總的角色。

我的回饋：

1. 簡報整體結構不錯，數據充分。但語氣需要調整——不要等客戶問才說壞消息，要主動坦承問題並立刻接改善措施。

2. Lisa 說的 before/after 比較圖一定要做，這是最有說服力的。

3. 10 月 88.5% 那個問題一定會被問，James 你的回答可以，但要更簡短——客戶不想聽太多解釋，他要聽你怎麼確保不會再發生。

4. Amy 對 NIC 問題的 Plan B 很好——有替代方案永遠比「我們在等對方」好。

5. Tom 的建議很好——把成本節省拆開來報，區分「正常成本節省」和「品質異常額外成本」。

6. Sarah 的 ROI 數字要放在投資頁面的顯眼位置。

修改行動項目："""),

        ("14:50", "James Chen", """謝謝 Robert 副總的回饋，非常實用。我整理一下修改行動項目：

1. James：修改 Executive Summary 的措辭，改用「挑戰與改善」框架 — 1/7 前
2. Lisa：交期頁面加入 before/after 比較圖 — 1/8 前
3. Amy：準備 NIC 問題的 Plan B talking points（含 Intel 替代驗證時程）— 1/8 前
4. Tom：成本節省頁面拆分為「正常節省」vs「品質異常成本」— 1/7 前
5. Sarah：投資提案頁面突出 ROI 275% 的數字 — 1/7 前
6. James：準備完整 Q&A 應對手冊（至少 15 個可能的問題+建議回答）— 1/10 前
7. James：1/12 再做一次 final dry run（30 分鐘，只走修改過的頁面）— 1/12

正式 QBR 是 1/15 下午 2 點，Dell 總部視訊。出席者確認：James（主講）、Lisa（交期）、Amy（品質）、Sarah（流程改善）。Tom 待命支援數據問題。

還有什麼遺漏的嗎？"""),

        ("14:53", "Lisa Wang", """建議加一個：我跟 Dell 的 procurement director Jennifer Walsh 比較熟，1/15 之前我先私下跟她聊一下，了解 Michael 這次特別關注什麼議題。這樣我們可以針對性準備。"""),

        ("14:54", "James Chen", """太好了，加入行動項目——Lisa 在 1/10 前跟 Jennifer 做非正式溝通，了解客戶的 pain points。

好的，今天的預演就到這裡。大家辛苦了，感謝 Robert 副總的嚴格質詢。剩下 10 天我們全力衝刺，1/15 見真章！散會。"""),
    ]

    for time, speaker, content in transcript:
        p = doc.add_paragraph()
        run_time = p.add_run(f"[{time}] ")
        run_time.bold = True
        run_time.font.color.rgb = RGBColor(128, 128, 128)
        run_time.font.size = Pt(10)
        run_speaker = p.add_run(f"{speaker}：")
        run_speaker.bold = True
        run_speaker.font.size = Pt(11)
        run_speaker.font.color.rgb = RGBColor(47, 84, 150)
        doc.add_paragraph(content.strip())

    doc.add_page_break()
    doc.add_heading('預演回饋摘要', level=2)
    feedback = [
        ("1", "簡報語氣", "主動坦承問題，不要等客戶問", "James", "1/7"),
        ("2", "交期頁面", "加入 before/after 比較圖", "Lisa", "1/8"),
        ("3", "NIC 問題", "準備 Plan B talking points（Intel 替代）", "Amy", "1/8"),
        ("4", "成本節省", "拆分為「正常節省」vs「品質異常成本」", "Tom", "1/7"),
        ("5", "投資提案", "突出 ROI 275%", "Sarah", "1/7"),
        ("6", "Q&A 手冊", "準備 15+ 個問題與建議回答", "James", "1/10"),
        ("7", "客戶探口風", "Lisa 私下與 Jennifer Walsh 溝通", "Lisa", "1/10"),
        ("8", "Final Dry Run", "再走一遍修改後的頁面（30 min）", "James", "1/12"),
    ]
    add_styled_table(doc, ["#", "項目", "修改內容", "負責人", "截止日"], feedback, "C55A11")

    path = os.path.join(OUTPUT_DIR, "英業達_QBR預演會議逐字稿.docx")
    doc.save(path)
    print(f"✅ Created: {path}")

if __name__ == "__main__":
    create_qbr_dryrun_transcript()
